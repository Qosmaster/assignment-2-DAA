"""Build the plain-English report and README from the measured CSV files.
The Java project runs without this optional Python script.
"""
from pathlib import Path
import csv
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
with (ROOT / 'results/tables/summary.csv').open() as source:
    ROWS = list(csv.DictReader(source))
LOOKUP = {(r['workload'], r['structure'], int(r['n'])): r for r in ROWS}
PAGES = []
BLOCKS = []


def page():
    global BLOCKS
    BLOCKS = []
    PAGES.append(BLOCKS)


def h(text, level=2):
    BLOCKS.append(('h', level, text))


def p(text):
    BLOCKS.append(('p', text))


def table(headers, rows, widths):
    BLOCKS.append(('table', headers, rows, widths))


def image(filename, caption, width=6.0):
    BLOCKS.append(('image', filename, caption, width))


def code(text):
    BLOCKS.append(('code', text))


def selected(workloads):
    rows = [r for r in ROWS if r['workload'] in workloads]
    return sorted(rows, key=lambda r: (int(r['n']), r['workload'], r['structure']))


def ms(workload, structure, n=100000):
    return f"{float(LOOKUP[(workload, structure, n)]['mean_ms']):.6f}"


def count(workload, structure, metric, n=100000):
    return f"{int(LOOKUP[(workload, structure, n)]['mean_' + metric]):,}"


def short_name(structure):
    if structure == 'DynamicArray':
        return 'Array'
    if structure == 'LinkedList':
        return 'List'
    return 'Heap'


def numbers(workload, metric, theory_array, theory_list):
    output = []
    for r in selected([workload]):
        theory = theory_array
        if r['structure'] == 'LinkedList':
            theory = theory_list
        output.append([f"{int(r['n']):,}", short_name(r['structure']),
                       r['mean_ms'], f"{int(r['mean_' + metric]):,}", theory])
    return output


page()
h('Assignment 2', 1)
p('Algorithmic Analysis, Correctness and Performance Trade-offs')
p('Student: ____________________    Group: ____________________')
h('1. Overview', 1)
p('This project compares a Dynamic Array, a Linked List, and a Min-Heap. The three structures are written in Java. They store integers. The aim is to explain how they work, prove two operations, and compare their speed using the four workloads in the assignment [1].')
table(['Structure', 'How it stores values', 'Required methods'], [
    ['Dynamic Array', 'An int array that doubles when full.', 'add(x), add(index, x), remove(index), get(index), contains(x)'],
    ['Linked List', 'Nodes. Each node stores a value and a link to the next node.', 'add(x), add(index, x), remove(index), get(index), contains(x)'],
    ['Min-Heap', 'An int array where every parent is <= its children.', 'insert(x), peekMin(), extractMin()']
], [1.10, 2.25, 3.20])
h('1.1 Main ideas in the code')
p('size is the number of stored values. Capacity is the length of the internal array. head is the first list node, and tail is the last. The array and heap start with space for 16 values. Neither one shrinks its array after removals.')
p('IntList is a small interface: a shared list of methods for the two list classes. It lets the same test code work with either class. Metrics stores three counters: accesses, comparisons, and movements. Demo.java shows small examples before the larger tests.')
p('The custom structures do not use ready-made collections as storage. Standard Java collections are used to check answers and to hold benchmark result records. The Java source has no streams, lambdas, recursion, or external libraries.')
h('1.2 Tests and error handling')
p('The tests check empty structures, one value, many values, duplicates, negative values, boundary positions, invalid indices, and 100,000-element inputs. They also check repeated array growth and list operations after the list becomes empty.')
p('Array results are compared with ArrayList; list results with java.util.LinkedList; heap results with PriorityQueue and a sorted array. Mixed heap tests check the heap rule after every insertion and removal. The latest run passed 188,698 checks. Its output is saved in results/test-results.txt [3].')
p('get and remove accept indices from 0 to size - 1. Insertion also accepts index size. Invalid indices throw IndexOutOfBoundsException. Reading or removing a minimum from an empty heap throws NoSuchElementException. These errors do not change the stored values.')

page()
h('2. Complexity Analysis', 1)
h('2.1 Meaning of the symbols')
p('n means the number of values before one operation. m means the number of operations in a workload. O gives an upper growth bound. Ω gives a lower bound. Θ gives a matching upper and lower bound. For example, Θ(n) means both O(n) and Ω(n). These symbols are not names for worst, best, and average cases.')
p('The analysis counts simple steps such as a comparison, an array access, or following a link. Average indexed operations assume a uniformly chosen valid index. Average search uses 50% present values and 50% absent values. Extra space means temporary working memory, not the stored data itself.')
h('2.2 Dynamic Array')
table(['Operation', 'Best', 'Average / sequence', 'Worst', 'Extra space'], [
    ['add(x)', 'Θ(1)', 'Θ(1) amortized*', 'Θ(n)', 'Θ(n) on growth; otherwise Θ(1)'],
    ['add(index, x)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(n) on growth; otherwise Θ(1)'],
    ['remove(index)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)'],
    ['get(index)', 'Θ(1)', 'Θ(1)', 'Θ(1)', 'Θ(1)'],
    ['contains(x)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)']
], [1.20, .55, 1.45, .55, 2.80])
p('get reads one position directly. contains checks values until it finds a match or reaches the end. Insertion shifts n - index values right. Removal shifts n - index - 1 values left. An operation at the end can avoid shifting. An operation near the front can move almost the whole array.')
p('*Amortized means the cost per operation over a whole sequence, not over random input values. Doubling copies 16, 32, 64, and so on. Over N appends, all these copies together take O(N) steps. The N new writes also take Ω(N). So N appends take Θ(N), or Θ(1) per append. One append to a full array still takes Θ(n).')
h('2.3 Linked List')
table(['Operation', 'Best', 'Average', 'Worst', 'Extra space'], [
    ['add(x)', 'Θ(1)', 'Θ(1)', 'Θ(1)', 'Θ(1)'],
    ['add(index, x)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)'],
    ['remove(index)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)'],
    ['get(index)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)'],
    ['contains(x)', 'Θ(1)', 'Θ(n)', 'Θ(n)', 'Θ(1)']
], [1.20, .55, 1.45, .55, 2.80])
p('add(x) uses tail, so it does not scan the list. Head insertion and removal are also constant time. Other indexed operations usually need to follow links first. get(i) visits i + 1 nodes. Search checks nodes one by one. Changing two links is cheap, but finding the correct node by index may take Θ(n) time. Removing the tail also needs its previous node.')

page()
h('2.4 Min-Heap')
table(['Operation', 'Best', 'Average / sequence', 'Worst'], [
    ['insert(x)', 'Θ(1)', 'Θ(1) expected amortized for random-order construction*', 'Θ(n) with growth; Θ(log n) without growth'],
    ['peekMin()', 'Θ(1)', 'Θ(1)', 'Θ(1)'],
    ['extractMin()', 'Θ(1)', 'Θ(log n), averaged over all removals of random distinct values*', 'Θ(log n)']
], [1.05, .60, 2.55, 2.35])
p('insert puts the new value at the end. It compares the value with its parent and swaps them when needed. Each swap moves one level up. A heap has about log2(n) levels, so the upward loop has a worst-case Θ(log n) cost. If the array is full, copying it adds Θ(n) time.')
p('peekMin returns the root at data[0]. The root is the minimum because every parent is <= its children. extractMin saves the root, moves the last value to the root, and moves it down by swapping with the smaller child. It checks at most two key comparisons per level. The worst case is Θ(log n); an immediate stop can take Θ(1).')
p('*The average insertion entry has a specific meaning. Inserting distinct values in random order into an empty heap takes expected Θ(n) total time [2]. With doubling, this gives expected amortized Θ(1) per insertion. It is not a guarantee for every insertion or every possible input. Descending input can take Θ(n log n) for all insertions. O(log n) amortized per insertion is a safe general bound.')
p('Removing every minimum puts the values in sorted order. For random distinct values, a full drain has expected Θ(n log n) comparisons, or Θ(log n) per extraction averaged over the drain. The upper bound follows from the heap height. The lower bound follows from comparison sorting: sorting random distinct values needs Ω(n log n) comparisons, and the expected build phase takes only O(n).')
h('2.5 Memory')
p('Heap insertion uses Θ(n) extra memory when a larger array is made; otherwise it uses Θ(1). peekMin and extractMin use Θ(1) extra memory. The code uses loops rather than recursive calls.')
p('The list stores Θ(n) nodes. The array and heap keep their allocated capacity after deletion. Therefore their retained memory follows their largest earlier size, not necessarily the current size. During growth, both the old and new arrays exist for a short time.')
h('2.6 One operation versus a whole workload')
p('For m random accesses, array work is Θ(m) and expected list work is Θ(mn). For m searches with half the values absent, both structures take expected Θ(mn). List head changes take Θ(m). List middle changes take Θ(mn), because each operation must reach the same fixed middle index.')
p('Repeated array insertion grows the structure. At a fixed index i, the shifts are m(n - i) + m(m - 1)/2, plus any growth copies. This gives Θ(mn + m²), not simply Θ(mn) when m also changes. In this assignment m is fixed at 1,000 for these changes.')

page()
h('3. Correctness', 1)
p('A loop invariant is a statement that stays true each time a loop begins. A complete proof checks the start, each step, and the end. The following proofs refer to the actual loops in the Java source.')
h('3.1 DynamicArray.add(index, value)')
p('Before insertion, size is s and index is between 0 and s. Call the original values A[0] to A[s - 1]. ensureCapacity keeps those values and makes room for one more. Counter updates are omitted from this short code extract; they do not change any stored value.')
code('for (int i = size; i > index; i--) {\n    data[i] = data[i - 1];\n}\ndata[index] = value;\nsize++;')
p('Invariant: at the start of a step with position i, index <= i <= s. Every position k below i still contains A[k]. Every position k above i, up to s, contains A[k - 1]. So the processed right part has moved one place right, and the unread left part is unchanged.')
p('Initialization: i starts at s. No values have moved yet. The unchanged part is the full original array, and the processed part is empty. The invariant is true.')
p('Maintenance: data[i - 1] is still the original value A[i - 1]. Copying it to data[i] puts that value in its correct new place. Decreasing i by one adds this position to the processed part. The remaining unread values are unchanged, so the invariant stays true.')
p('Termination: i - index decreases by one and cannot go below zero during the loop. The loop stops at i = index. All original values from index onward have moved right once. Earlier values are unchanged. Writing the new value at index and increasing size creates the required sequence without losing or reordering the old values. End insertion is also covered: its loop runs zero times.')
h('3.2 LinkedList.contains(value)')
p('Precondition: the list is valid and has a finite chain of nodes. current starts at head. The loop reads current.value, returns true on a match, and otherwise moves to current.next.')
p('Invariant: every node before current has already been checked and does not contain the target. current is the next node to check, or null when there are no nodes left. No value or link has changed.')
p('Initialization: current is head. There are no earlier nodes, so the invariant is true. It also holds for an empty list, where head is null.')
p('Maintenance: a match proves the value exists, so returning true is correct. Without a match, moving to the next node adds one known non-matching node to the checked part. The invariant stays true.')
p('Termination: each unsuccessful step leaves one fewer node to check. Since the list is finite and valid, the loop either returns true or reaches null. At null, every node has been checked and none matched, so returning false is correct. Thus the method returns true exactly when the target is present, including lists with duplicates.')

page()
h('4. Experimental Setup', 1)
p('The four input sizes are 100, 1,000, 10,000, and 100,000. Every experiment is repeated five times. The reported time is the arithmetic mean: add the five times and divide by five. System.nanoTime() measures the operation loops. Results are then converted to milliseconds.')
table(['Workload', 'Structures', 'Operations per experiment'], [
    ['1. Random access', 'Array and list', '10,000 get(index) calls'],
    ['2. Search', 'Array and list', '1,000 contains(value) calls'],
    ['3. Insert / remove', 'Array and list', '1,000 separate insertions or removals at 0 or fixed n / 2'],
    ['4. Priority processing', 'Heap', 'n insertions into an empty heap, then n extractions']
], [1.25, 1.25, 4.05])
p('Random(42) is used for each input size. Values 0 to n - 1 are shuffled into a random order. They are distinct, not independently sampled with replacement. Access indices, 1,000 new insertion values, and search values are prepared before timing. Exactly 500 search values are present and 500 are absent. Both structures and all five runs use the same prepared input.')
p('Input generation, list filling, printing, restoration, and answer checks are outside the timer. Array growth, new list nodes, and counter updates caused by the measured operations are inside it. Two unrecorded warm-up rounds run at n = 1,000 first. The order of the two list structures changes between repetitions.')
h('4.1 What the counters mean')
p('Array accesses count slot reads and writes. Moving one existing value counts one movement and two accesses. Growth copies are included. List accesses count visited or directly handled nodes, including a new or removed node. List movements are zero because existing values are not shifted. These are different units, not equal CPU instructions.')
p('Comparisons count value checks in contains or comparisons between heap keys. Index checks and loop conditions are not included. Heap comparisons are counted, but heap accesses are not: a zero access field for the heap means unused, not no memory work. All counters use long.')
h('4.2 Workload 3: the small-size problem')
p('The brief asks for 1,000 removals after restoring n original values [1]. For n = 100, this is impossible in one list. With the fixed middle index, n = 1,000 also allows only 500 removals before that index becomes invalid.')
p('The code removes as many values as are valid, restores the original structure outside the timer, and continues until 1,000 successful removals are complete. It adds only the measured removal times and counts. The index stays 0 or the original n / 2. This interpretation needs the instructor\'s approval; it is not silently treated as an exact instruction from the brief.')
table(['n', 'Front batches', 'Middle index', 'Middle batches'], [
    ['100', '10', '50', '20'], ['1,000', '1', '500', '2'],
    ['10,000', '1', '5,000', '1'], ['100,000', '1', '50,000', '1']
], [1.15, 1.60, 1.90, 1.90])
p('A batch includes one prepared structure. All insertions start from a separate original structure. Measurements were made in a shared Linux container with OpenJDK 21, not on the student\'s computer. Full version and memory details are in results/environment.txt. Times may vary because of Java warm-up, scheduling, and memory use [3].')

page()
h('5. Results', 1)
p('The files contain 280 measured runs and 56 averages. Each table row is the total time for one workload, averaged over five runs. All values below come from the current CSV files [3]. Array and List mean the custom Java classes.')
h('5.1 Workload 1: random access')
table(['n', 'Type', 'Mean ms', 'Accesses', 'Total time bound'],
      numbers('W1_access', 'accesses', 'Θ(m)', 'Θ(mn), expected'),
      [1.00, .60, 1.10, 1.60, 2.25])
image('01_access_time.png', 'Figure 1. Time for 10,000 random accesses. Bars show the smallest and largest measured times.')
p(f"At n = 100,000, the array takes {ms('W1_access', 'DynamicArray')} ms and the list takes {ms('W1_access', 'LinkedList')} ms. The array always makes 10,000 slot accesses. The list makes {count('W1_access', 'LinkedList', 'accesses')} node visits at the largest size.")
p('The array jumps directly to an index. The list starts from head each time. A random list index needs about (n + 1) / 2 node visits on average. The counts therefore agree with Θ(1) array access and expected Θ(n) list access. Small timing changes do not change these bounds.')

page()
h('5.2 Workload 2: search', 1)
p('m = 1,000. Half of the search values are present and half are absent. Both structures hold values in the same order and receive the same queries.')
table(['n', 'Type', 'Mean ms', 'Comparisons', 'Total time bound'],
      numbers('W2_search', 'comparisons', 'Θ(mn), expected', 'Θ(mn), expected'),
      [1.00, .60, 1.10, 1.60, 2.25])
image('03_search_time.png', 'Figure 2. Search time for 1,000 queries. Bars show the measured minimum and maximum.')
p(f"At n = 100,000, both structures make {count('W2_search', 'DynamicArray', 'comparisons')} comparisons. The array takes {ms('W2_search', 'DynamicArray')} ms; the list takes {ms('W2_search', 'LinkedList')} ms.")
p('A successful search uses about (n + 1) / 2 comparisons on average. A failed search uses n. This explains the near-linear growth in the count when n increases. Equal comparison counts do not mean equal time: following next links and reading adjacent array positions have different costs.')

page()
h('5.3 Workload 3A: insertion', 1)
p('m = 1,000. Front means index 0. Middle means the fixed original n / 2. The metric is moved values for the array and handled nodes for the list. Array growth copies are included.')
rows = []
for r in selected(['W3_insert_front', 'W3_insert_middle']):
    front = r['workload'].endswith('front')
    position = 'Front' if front else 'Middle'
    metric = 'mean_movements' if r['structure'] == 'DynamicArray' else 'mean_accesses'
    theory = 'Θ(mn + m²)' if r['structure'] == 'DynamicArray' else ('Θ(m)' if front else 'Θ(mn)')
    rows.append([f"{int(r['n']):,}", position, short_name(r['structure']),
                 r['mean_ms'], f"{int(r[metric]):,}", theory])
table(['n', 'Position', 'Type', 'Mean ms', 'Metric', 'Total bound'], rows,
      [.80, .75, .50, 1.00, 1.35, 2.15])
image('05_insert_time.png', 'Figure 3. Time for 1,000 insertions at the front or fixed middle index.')
p('The array shifts values to make space. Head insertion in the list only changes links, so it is cheaper in number of steps. Middle insertion in the list still needs a walk from head. The array grows during the 1,000 insertions, which is why its full bound includes m².')

page()
h('5.4 Workload 3B: removal', 1)
p('Each row has 1,000 successful removals. Small structures are restored outside timing as explained in Section 4.2. The counter units are the same as in the insertion table.')
rows = []
for r in selected(['W3_remove_front', 'W3_remove_middle']):
    front = r['workload'].endswith('front')
    position = 'Front' if front else 'Middle'
    metric = 'mean_movements' if r['structure'] == 'DynamicArray' else 'mean_accesses'
    theory = 'Θ(mn)' if r['structure'] == 'DynamicArray' or not front else 'Θ(m)'
    rows.append([f"{int(r['n']):,}", position, short_name(r['structure']),
                 r['mean_ms'], f"{int(r[metric]):,}", r['batches'], theory])
table(['n', 'Position', 'Type', 'Mean ms', 'Metric', 'Batches', 'Total bound'], rows,
      [.75, .72, .48, .93, 1.20, .82, 1.65])
image('06_remove_time.png', 'Figure 4. Time for 1,000 removals. Restoration is not timed.')
p('Front removal shifts almost all remaining array values. The list can just move head to the next node. In the middle, the array shifts fewer values, but the list still needs to find the previous node. The array bound shown here uses the documented batch rule and the tested sizes.')

page()
h('5.5 Workload 4: priority processing', 1)
p('The heap starts empty. It receives n values, then returns the minimum n times. These two parts have separate timers and comparison counters. The returned sequence is compared with a sorted copy of the input outside the timer.')
rows = []
for r in selected(['W4_insert', 'W4_extract']):
    inserting = r['workload'] == 'W4_insert'
    rows.append([f"{int(r['n']):,}", 'Insert' if inserting else 'Extract', r['mean_ms'],
                 f"{int(r['mean_comparisons']):,}", 'Θ(n), expected' if inserting else 'Θ(n log n), expected'])
table(['n', 'Phase', 'Mean ms', 'Comparisons', 'Total time bound'], rows,
      [.80, .75, 1.10, 1.50, 2.40])
image('07_heap_time.png', 'Figure 5. Separate times for inserting n values and extracting n minima.')
p(f"At n = 100,000, insertion makes {count('W4_insert', 'MinHeap', 'comparisons')} comparisons and extraction makes {count('W4_extract', 'MinHeap', 'comparisons')}. The mean times are {ms('W4_insert', 'MinHeap')} ms and {ms('W4_extract', 'MinHeap')} ms.")
p('Random-order insertion often stops after only a few upward comparisons. Repeated extraction usually moves values farther down. This agrees with the random-order averages in Section 2.4. Both phases have O(n log n) worst-case total bounds, including array growth. All extracted sequences passed the non-decreasing-order check.')
p('peekMin is tested and analyzed, but it is not separately timed because this workload asks for insertion and extraction times only.')

page()
h('5.6 Operation-count graphs', 1)
p('These graphs show work counts rather than speed. Logarithmic axes keep small and large values readable. The exact numbers are in the tables and CSV files.')
image('02_access_counts.png', 'Figure 6. Array accesses stay fixed; list node visits grow with n.', 5.45)
image('04_search_comparisons.png', 'Figure 7. The two structures have identical search comparison counts.', 5.45)
image('08_heap_comparisons.png', 'Figure 8. Heap insertion and extraction comparison counts.', 5.45)
p('Counts are identical across the five runs because the input is fixed. verify_results.py checks averages, repeat counts, search hits, and exact insertion/removal counts. These checks passed for all 56 experiments [3].')

page()
h('6. Discussion', 1)
h('6.1 How does increasing n affect the workloads?')
p('Array access keeps the same number of steps. List access and both searches do more work. Array front changes shift more values. List head changes stay cheap. Middle list changes need more link visits. Heap extraction has more levels to check.')
h('6.2 Which results agree with the theory?')
p('Array access always counts 10,000 reads. List head changes always count 1,000 handled nodes. List middle changes count 1,000(n / 2 + 1) nodes. Search counts and array movement counts follow their expected formulas. Heap comparison growth matches the random-order analysis.')
h('6.3 Where do time results differ from a simple prediction?')
a100 = ms('W1_access', 'DynamicArray', 100)
a100k = ms('W1_access', 'DynamicArray', 100000)
p(f'Array access has a mean of {a100} ms at n = 100 and {a100k} ms at n = 100,000, even though the access count is equal. Small timings can vary because of Java compilation, scheduling, and memory effects. These causes were not measured separately. None of the observed runs was removed.')
h('6.4 Why can the same Big-O give different times?')
p('Big-O does not show every small cost. Both searches are linear, but an array reads nearby slots while a list follows links. The tables show equal comparisons but different times. This does not contradict the analysis.')
h('6.5 What implementation details matter?')
p('tail makes list append constant time. Doubling makes most array appends cheap but creates occasional long copies. List insertion creates a node. The counters also add work inside the timer, so these measurements describe this instrumented code.')
h('6.6 When is a Dynamic Array useful?')
p('It is useful when indexed reads are common. It gives direct access and cheap amortized appends. It was also faster than the list for search and for large middle-index changes in this run. Repeated front changes are its main weakness here.')
h('6.7 When is a Linked List useful?')
p('It is useful for frequent head insertion and removal. Linking at a known node is also cheap. However, this project takes numeric indices, so it must first find an interior node. List insertion is not always constant time.')
h('6.8 Why use a heap for priority processing?')
p('The minimum is at the root, and removal fixes one path. Removing all minima has O(n log n) worst-case work. Repeatedly scanning an unsorted collection would take Θ(n²) checks in total. A heap is not a fully sorted array.')
h('6.9 How does the workload affect the choice?')
p('Choose for the operations that happen most often, not for one attractive complexity value. Indexed reading, head changes, and next-minimum processing need different strengths. These tests use one seed and four sizes; they support the analysis but do not prove speed for every program or computer.')

page()
h('7. Design Recommendations', 1)
table(['Main task', 'Structure', 'Reason'], [
    ['Read by index', 'Dynamic Array', 'Direct access without following links.'],
    ['Search an unsorted sequence', 'Dynamic Array here', 'Both are linear; the array was faster in these runs.'],
    ['Insert/remove at the front', 'Linked List', 'Change head without shifting other values.'],
    ['Change a middle numeric index', 'Dynamic Array here', 'Both are linear; list traversal was slower at larger sizes.'],
    ['Repeatedly take the minimum', 'Min-Heap', 'Direct minimum lookup and short repair paths.']
], [2.10, 1.35, 3.10])
h('8. Conclusion', 1)
p('No structure is best for every job. The Dynamic Array is strong for indexed access. The Linked List is strong for head changes. The Min-Heap is useful for priority processing. Counts agree with the main complexity bounds, while measured times also depend on small implementation and computer effects.')
p('All required methods and tests are included. The two proofs cover indexed array insertion and list search. The Workload 3 removal rule still needs the instructor\'s approval. The GitHub repository also needs to be published before submission.')
h('Appendix: running and submitting')
p('On Windows, run check.bat first. It compiles the source, runs Demo, and runs Tests. It does not replace the saved measurements. A JDK is needed; this project was tested with JDK 21. On macOS/Linux use sh check.sh.')
p('run.bat or sh run.sh starts the full benchmark and replaces the result CSV files. Run it in a copy of the project to keep the supplied results unchanged. To rebuild all graphs and documents after new measurements, use the optional Python scripts:')
code('python -m pip install matplotlib python-docx\npython scripts/verify_results.py\npython scripts/make_plots.py\npython scripts/build_report.py')
p('The brief asks for tables, graphs, an individual report, and a GitHub repository link. It does not separately require screenshots [1]. A screenshot of your own successful test run can be added as extra evidence, but it does not replace the source files or results.')
p('The ZIP includes the existing local development history and the actual simplification changes. These assistant-prepared commits are not past work done by the student. Publish honestly under your authorized account; see docs/PUBLISHING.md. Review the code and fill in the name and group before submission.')
h('Sources', 2)
p('[1] Supplied Assignment 2-1.pdf, Sections 3-14. Source of the requirements.')
p('[2] Bollobás, B., and Simon, I. (1985). Repeated random insertion into a priority queue. Journal of Algorithms 6(4), 466-477. DOI: 10.1016/0196-6774(85)90028-8. Source of the random-order expected heap construction result.')
p('[3] This project: src/*.java, results/tables/raw.csv, summary.csv, test-results.txt, benchmark-log.txt, verification.txt, and environment.txt. Source of the tests, measurements, and checks. Proofs and workload explanations follow the supplied code.')


def markdown(block):
    kind = block[0]
    if kind == 'h':
        return '#' * (block[1] + 1) + ' ' + block[2] + '\n\n'
    if kind == 'p':
        return block[1] + '\n\n'
    if kind == 'code':
        return '```\n' + block[1] + '\n```\n\n'
    if kind == 'image':
        return f'![{block[2]}](results/plots/{block[1]})\n\n'
    if kind == 'table':
        rows = [block[1], ['---'] * len(block[1])] + block[2]
        output = []
        for row in rows:
            output.append('| ' + ' | '.join(str(x).replace('|', '/') for x in row) + ' |')
        return '\n'.join(output) + '\n\n'
    raise ValueError(kind)


readme = '# Assignment 2: Java Data Structures\n\n'
readme += 'Start with `START_HERE.md`. Use `check.bat` on Windows for a safe test run.\n\n'
for blocks in PAGES:
    for block in blocks:
        readme += markdown(block)
(ROOT / 'README.md').write_text(readme, encoding='utf-8')
results_text = '# Measured results\n\n'
for blocks in PAGES[5:11]:
    for block in blocks:
        if block[0] != 'image':
            results_text += markdown(block)
(ROOT / 'results/tables/results.md').write_text(results_text, encoding='utf-8')
for number in (1, 2, 3, 4):
    with (ROOT / f'results/tables/workload_{number}.csv').open('w', newline='') as destination:
        writer = csv.DictWriter(destination, fieldnames=list(ROWS[0]))
        writer.writeheader()
        writer.writerows(r for r in ROWS if r['workload'].startswith(f'W{number}_'))

# Simple Word layout, without a separate cover page.
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.2677)
section.page_height = Inches(11.6929)
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.82)
section.right_margin = Inches(.82)
section.header_distance = Inches(.25)
section.footer_distance = Inches(.25)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 11)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True
for name in ['Header', 'Footer']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.size = Pt(8)
section.header.paragraphs[0].text = 'Assignment 2 | Java Data Structures'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Assignment 2 - Java Data Structures'
doc.core_properties.author = 'Assignment Project Builder'
doc.core_properties.subject = 'Plain-English report, proofs, and measured results'


def add_table(block):
    headers, rows, widths = block[1:]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for j, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            if j == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8E8E8')
                cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(3)
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(9)
                    run.bold = j == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)


for number, blocks in enumerate(PAGES):
    if number > 0:
        doc.add_page_break()
    for block in blocks:
        if block[0] == 'h':
            doc.add_heading(block[2], level=block[1])
        elif block[0] == 'p':
            doc.add_paragraph(block[1])
        elif block[0] == 'table':
            add_table(block)
        elif block[0] == 'code':
            para = doc.add_paragraph()
            para.paragraph_format.keep_together = True
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(block[1])
            run.font.name = 'Liberation Mono'
            run.font.size = Pt(9)
        elif block[0] == 'image':
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_after = Pt(2)
            shape = para.add_run().add_picture(str(ROOT / 'results/plots' / block[1]), width=Inches(block[3]))
            shape._inline.docPr.set('descr', block[2])
            caption = doc.add_paragraph(block[2])
            caption.paragraph_format.space_after = Pt(5)
            for run in caption.runs:
                run.font.size = Pt(8.5)
                run.italic = True
path = ROOT / 'report/Assignment_2_Report.docx'
path.parent.mkdir(exist_ok=True)
doc.save(path)
print('Created the plain-English Word report, matching README, and result tables.')
print(f'Planned pages: {len(PAGES)}. Render the Word file to check the final page layout.')
